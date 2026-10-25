from pathlib import Path
import re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
B=Path(__file__).resolve().parent;ROOT=B.parent
OUT=ROOT/'output'/'第6课_从隋唐盛世到五代十国_教案.docx'
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(1.8)
for name,size,font in [('Normal',11,'宋体'),('Title',20,'黑体'),('Heading 1',15,'黑体'),('Heading 2',12,'黑体'),('Heading 3',11,'黑体')]:
 st=d.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font);st.paragraph_format.line_spacing=1.2;st.paragraph_format.space_after=Pt(5)
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(9);st.paragraph_format.keep_with_next=True
d.styles['Title'].paragraph_format.space_after=Pt(14)
def inline(p,t):
 for a in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))',t):
  if not a:continue
  m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',a)
  if m:
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:eastAsia'),'宋体');pr.append(f);c=OxmlElement('w:color');c.set(qn('w:val'),'245B75');pr.append(c);r.append(pr);tt=OxmlElement('w:t');tt.text=m[1];r.append(tt);h.append(r);p._p.append(h)
  else:
   r=p.add_run(a[2:-2] if a.startswith('**') else a);r.bold=a.startswith('**')
lines=(B/'lesson.md').read_text(encoding='utf8').splitlines();i=0
while i<len(lines):
 t=lines[i].strip();i+=1
 if not t:continue
 if t.startswith('|'):
  rows=[t]
  while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
  rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)];data=[[c.strip() for c in r.strip('|').split('|')] for r in rows];n=len(data[0]);tab=d.add_table(rows=0,cols=n);tab.autofit=False;tab.alignment=WD_TABLE_ALIGNMENT.CENTER
  widths=([6.4,11] if len(data)==5 else [3.4,14]) if n==2 else [2.5,1.7,7.3,5.9]
  for col,w in zip(tab.columns,widths):col.width=Cm(w)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  tab._tbl.tblPr.append(borders)
  for ri,row in enumerate(data):
   cells=tab.add_row().cells
   trpr=tab.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if ri==0:trpr.append(OxmlElement('w:tblHeader'))
   for ci,value in enumerate(row):
    cell=cells[ci];cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     a=OxmlElement('w:'+edge);a.set(qn('w:w'),'90');a.set(qn('w:type'),'dxa');marg.append(a)
    pr.append(marg)
    if ri==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EBEF');pr.append(shade)
    p=cell.paragraphs[0];inline(p,value);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
    if ri==0 or (n==4 and ci==1):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
  d.add_paragraph().paragraph_format.space_after=Pt(0)
 elif t.startswith('# '):p=d.add_paragraph(t[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif t.startswith('### '):d.add_paragraph(t[4:],'Heading 2')
 elif t.startswith('## '):d.add_paragraph(t[3:],'Heading 1')
 else:
  p=d.add_paragraph();inline(p,('• '+t[2:]) if t.startswith('- ') else t)
  if re.fullmatch(r'\*\*\d+\..*\*\*',t):p.paragraph_format.keep_with_next=True
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field);foot.add_run(' 页')
for r in foot.runs:r.font.size=Pt(9)
d.core_properties.title='第6课《从隋唐盛世到五代十国》教案';d.core_properties.author='';d.core_properties.subject='高中历史 40分钟课堂教学设计'
d.save(OUT);print(OUT)

